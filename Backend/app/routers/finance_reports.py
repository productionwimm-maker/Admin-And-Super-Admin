"""Finance & reporting — SUPER-ADMIN ONLY.

A single place to see every rupee moving through WIMM and to export accounting
sheets. Data is read live from Firestore:

  • subscription_payments  → platform REVENUE (pharmacies paying to be listed)
  • orders (paid)          → GMV flowing through + delivery fees collected
  • refunds                → money paid back out (a cost)

Money model (platform's own P&L, not the pharmacies'):
  revenue   = subscription income
  gmv       = paid-order value (passes through to pharmacies, shown for context)
  delivery  = delivery fees (customer-funded, pass-through to the courier)
  expenses  = refunds + estimated gateway fees (2% of revenue)
  gst       = 18% component of revenue
  profit    = revenue − expenses           (negative ⇒ loss)

Everything can be sliced by hour / day / week / month / financial-year (India
FY = 1 Apr–31 Mar) and exported as neatly formatted Excel.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse

from .. import store
from ..security import CurrentAdmin, require_superadmin

router = APIRouter(prefix="/api/finance", tags=["finance-reports"])

GST_RATE = 18          # %
GATEWAY_FEE_PCT = 2    # % of revenue, estimated
IST = timezone(timedelta(hours=5, minutes=30))


# ── helpers ────────────────────────────────────────────────────────────────
def _ms(dt: datetime) -> int:
    return int(dt.timestamp() * 1000)


def _now() -> datetime:
    return datetime.now(IST)


def _parse_window(period: Optional[str], frm: Optional[int], to: Optional[int]) -> tuple[int, int]:
    """Resolve a [from_ms, to_ms) window from an explicit range or a named period."""
    if frm is not None and to is not None:
        return frm, to
    now = _now()
    p = (period or "month").lower()
    if p == "hour":
        start = now.replace(minute=0, second=0, microsecond=0)
    elif p == "day":
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    elif p == "week":
        d = now.replace(hour=0, minute=0, second=0, microsecond=0)
        start = d - timedelta(days=d.weekday())
    elif p == "year" or p == "fy":
        # Indian FY starts 1 Apr.
        y = now.year if now.month >= 4 else now.year - 1
        start = datetime(y, 4, 1, tzinfo=IST)
    else:  # month
        start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    return _ms(start), _ms(now) + 1


def _rows(from_ms: int, to_ms: int) -> dict:
    """Pull and window the raw transaction rows once."""
    def win(rows, field="createdAtMillis"):
        return [r for r in rows if from_ms <= int(r.get(field, 0) or 0) < to_ms]

    subs = win(store.collection("subscription_payments").list())
    orders_all = store.collection("orders").list()
    orders = win([o for o in orders_all if o.get("paymentStatus") == "PAID"])
    refunds = win(store.collection("refunds").list())
    return {"subs": subs, "orders": orders, "refunds": refunds}


def _totals(rows: dict) -> dict:
    revenue = sum(int(s.get("amountRupees") or 0) for s in rows["subs"])
    gmv = sum(int(o.get("totalAmount") or 0) for o in rows["orders"])
    delivery = sum(int(o.get("deliveryFee") or 0) for o in rows["orders"])
    refunds = sum(int(r.get("amount") or 0) for r in rows["refunds"])
    gateway = round(revenue * GATEWAY_FEE_PCT / 100)
    gst = round(revenue * GST_RATE / (100 + GST_RATE))
    expenses = refunds + gateway
    profit = revenue - expenses
    return {
        "revenue": revenue, "gmv": gmv, "delivery": delivery,
        "refunds": refunds, "gatewayFees": gateway, "gst": gst,
        "taxableValue": revenue - gst, "expenses": expenses,
        "profit": profit, "loss": (-profit if profit < 0 else 0),
        "subscriptionCount": len(rows["subs"]),
        "paidOrderCount": len(rows["orders"]),
        "refundCount": len(rows["refunds"]),
        "currency": "INR",
    }


def _bucket_key(ms: int, granularity: str) -> str:
    dt = datetime.fromtimestamp(ms / 1000, IST)
    g = granularity.lower()
    if g == "hour":
        return dt.strftime("%Y-%m-%d %H:00")
    if g == "week":
        d = dt - timedelta(days=dt.weekday())
        return d.strftime("W%Y-%m-%d")
    if g == "month":
        return dt.strftime("%Y-%m")
    if g in ("year", "fy"):
        y = dt.year if dt.month >= 4 else dt.year - 1
        return f"FY{y}-{str(y + 1)[2:]}"
    return dt.strftime("%Y-%m-%d")  # day


def _timeseries(rows: dict, granularity: str) -> list[dict]:
    buckets: dict[str, dict] = {}

    def b(ms):
        k = _bucket_key(int(ms or 0), granularity)
        return buckets.setdefault(k, {"bucket": k, "revenue": 0, "gmv": 0, "delivery": 0, "refunds": 0})

    for s in rows["subs"]:
        b(s.get("createdAtMillis"))["revenue"] += int(s.get("amountRupees") or 0)
    for o in rows["orders"]:
        d = b(o.get("createdAtMillis")); d["gmv"] += int(o.get("totalAmount") or 0); d["delivery"] += int(o.get("deliveryFee") or 0)
    for r in rows["refunds"]:
        b(r.get("createdAtMillis"))["refunds"] += int(r.get("amount") or 0)
    out = sorted(buckets.values(), key=lambda x: x["bucket"])
    for d in out:
        gateway = round(d["revenue"] * GATEWAY_FEE_PCT / 100)
        d["profit"] = d["revenue"] - d["refunds"] - gateway
    return out


# ── JSON endpoints (Money & Transactions sub-tab) ──────────────────────────
@router.get("/summary")
def summary(period: Optional[str] = "month", frm: Optional[int] = Query(None, alias="from"),
            to: Optional[int] = None, admin: CurrentAdmin = Depends(require_superadmin)):
    a, b = _parse_window(period, frm, to)
    t = _totals(_rows(a, b))
    t.update({"from": a, "to": b, "period": period})
    return t


@router.get("/breakdown")
def breakdown(period: Optional[str] = "month", frm: Optional[int] = Query(None, alias="from"),
              to: Optional[int] = None, admin: CurrentAdmin = Depends(require_superadmin)):
    a, b = _parse_window(period, frm, to)
    t = _totals(_rows(a, b))
    # Pie: where the money comes from (positive inflows).
    return {"slices": [
        {"name": "Subscriptions", "value": t["revenue"]},
        {"name": "Order GMV", "value": t["gmv"]},
        {"name": "Delivery", "value": t["delivery"]},
    ], "outflows": [
        {"name": "Refunds", "value": t["refunds"]},
        {"name": "Gateway fees", "value": t["gatewayFees"]},
    ]}


@router.get("/timeseries")
def timeseries(granularity: str = "day", period: Optional[str] = None,
               frm: Optional[int] = Query(None, alias="from"), to: Optional[int] = None,
               admin: CurrentAdmin = Depends(require_superadmin)):
    # Default window scales with granularity when none is given.
    if frm is None and to is None and period is None:
        period = {"hour": "day", "day": "month", "week": "year", "month": "fy", "fy": "fy"}.get(granularity, "month")
    a, b = _parse_window(period, frm, to)
    return {"granularity": granularity, "points": _timeseries(_rows(a, b), granularity)}


@router.get("/transactions")
def transactions(period: Optional[str] = "month", frm: Optional[int] = Query(None, alias="from"),
                 to: Optional[int] = None, limit: int = 200,
                 admin: CurrentAdmin = Depends(require_superadmin)):
    a, b = _parse_window(period, frm, to)
    rows = _rows(a, b)
    tx = []
    for s in rows["subs"]:
        tx.append({"type": "Subscription", "direction": "in", "amount": int(s.get("amountRupees") or 0),
                   "party": s.get("pharmacyPhone", ""), "ref": s.get("paymentId", ""),
                   "at": int(s.get("createdAtMillis", 0) or 0)})
    for o in rows["orders"]:
        tx.append({"type": "Order", "direction": "in", "amount": int(o.get("totalAmount") or 0),
                   "party": o.get("pharmacyName") or o.get("pharmacyPhone", ""), "ref": o.get("paymentId", ""),
                   "at": int(o.get("createdAtMillis", 0) or 0)})
    for r in rows["refunds"]:
        tx.append({"type": "Refund", "direction": "out", "amount": int(r.get("amount") or 0),
                   "party": r.get("orderId", ""), "ref": r.get("gatewayRef", ""),
                   "at": int(r.get("createdAtMillis", 0) or 0)})
    tx.sort(key=lambda x: x["at"], reverse=True)
    return {"count": len(tx), "transactions": tx[:limit]}


# ── Excel export (Reports sub-tab) ─────────────────────────────────────────
def _xlsx(export_type: str, granularity: str, rows: dict, window: tuple[int, int]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    HEAD = PatternFill("solid", fgColor="0E7C66")
    HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
    TITLE_FONT = Font(bold=True, size=14, color="0E7C66")
    MONEY = '#,##0" ₹"'
    thin = Side(style="thin", color="D0D7DE")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    CENTER = Alignment(horizontal="center", vertical="center")

    wb = Workbook()
    t = _totals(rows)
    a, b = window
    span = f"{datetime.fromtimestamp(a/1000, IST):%d %b %Y %H:%M} — {datetime.fromtimestamp(b/1000, IST):%d %b %Y %H:%M}"

    def style_header(ws, ncols, row=1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = HEAD; cell.font = HEAD_FONT; cell.alignment = CENTER; cell.border = BORDER

    def autofit(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def kv_sheet(ws, title, pairs):
        ws["A1"] = title; ws["A1"].font = TITLE_FONT
        ws["A2"] = span; ws["A2"].font = Font(italic=True, color="57606A")
        ws.append([]); ws.append(["Metric", "Amount (₹)"]); style_header(ws, 2, row=4)
        for k, v in pairs:
            ws.append([k, v])
            ws.cell(row=ws.max_row, column=2).number_format = MONEY
            for c in (1, 2):
                ws.cell(row=ws.max_row, column=c).border = BORDER
        autofit(ws, [34, 20])

    def series_sheet(ws, title):
        ws["A1"] = title; ws["A1"].font = TITLE_FONT
        ws.append([]); ws.append(["Period", "Revenue", "GMV", "Delivery", "Refunds", "Profit"])
        style_header(ws, 6, row=3)
        for d in _timeseries(rows, granularity):
            ws.append([d["bucket"], d["revenue"], d["gmv"], d["delivery"], d["refunds"], d["profit"]])
            for c in range(2, 7):
                ws.cell(row=ws.max_row, column=c).number_format = MONEY
            for c in range(1, 7):
                ws.cell(row=ws.max_row, column=c).border = BORDER
        autofit(ws, [18, 16, 16, 14, 14, 16])

    def tx_sheet(ws):
        ws["A1"] = "Transactions"; ws["A1"].font = TITLE_FONT
        ws.append([]); ws.append(["Date", "Type", "Direction", "Party", "Amount (₹)", "Reference"])
        style_header(ws, 6, row=3)
        allrows = []
        for s in rows["subs"]:
            allrows.append((int(s.get("createdAtMillis", 0) or 0), "Subscription", "IN", s.get("pharmacyPhone", ""), int(s.get("amountRupees") or 0), s.get("paymentId", "")))
        for o in rows["orders"]:
            allrows.append((int(o.get("createdAtMillis", 0) or 0), "Order", "IN", o.get("pharmacyName") or o.get("pharmacyPhone", ""), int(o.get("totalAmount") or 0), o.get("paymentId", "")))
        for r in rows["refunds"]:
            allrows.append((int(r.get("createdAtMillis", 0) or 0), "Refund", "OUT", r.get("orderId", ""), int(r.get("amount") or 0), r.get("gatewayRef", "")))
        allrows.sort(key=lambda x: x[0], reverse=True)
        for at, ty, di, party, amt, ref in allrows:
            ws.append([datetime.fromtimestamp(at/1000, IST).strftime("%Y-%m-%d %H:%M"), ty, di, party, amt, ref])
            ws.cell(row=ws.max_row, column=5).number_format = MONEY
            for c in range(1, 7):
                ws.cell(row=ws.max_row, column=c).border = BORDER
        autofit(ws, [18, 14, 11, 26, 14, 26])

    sheets_needed = {
        "budget": ["budget", "series"], "profit": ["profit"], "loss": ["loss"],
        "tax": ["tax"], "gst": ["gst"],
        "all": ["budget", "profit", "loss", "tax", "gst", "series", "transactions"],
    }.get(export_type, ["budget"])

    first = True
    for name in sheets_needed:
        ws = wb.active if first else wb.create_sheet()
        first = False
        if name == "budget":
            ws.title = "Budget"
            kv_sheet(ws, "Budget Summary", [
                ("Revenue (subscriptions)", t["revenue"]), ("Order GMV", t["gmv"]),
                ("Delivery collected", t["delivery"]), ("Refunds", t["refunds"]),
                ("Gateway fees (est.)", t["gatewayFees"]), ("Expenses", t["expenses"]),
                ("Net profit", t["profit"]), ("Loss", t["loss"]),
            ])
        elif name == "profit":
            ws.title = "Profit"
            kv_sheet(ws, "Profit & Loss", [
                ("Revenue", t["revenue"]), ("Less: Refunds", -t["refunds"]),
                ("Less: Gateway fees", -t["gatewayFees"]), ("Net profit", t["profit"]),
            ])
        elif name == "loss":
            ws.title = "Loss"
            kv_sheet(ws, "Loss Statement", [
                ("Refunds paid", t["refunds"]), ("Gateway fees", t["gatewayFees"]),
                ("Total outflow", t["expenses"]), ("Net (negative = loss)", t["profit"]),
                ("Loss", t["loss"]),
            ])
        elif name == "tax":
            ws.title = "Tax"
            kv_sheet(ws, "Tax Summary", [
                ("Gross revenue", t["revenue"]), ("Taxable value", t["taxableValue"]),
                (f"GST ({GST_RATE}%, incl.)", t["gst"]),
            ])
        elif name == "gst":
            ws.title = "GST"
            kv_sheet(ws, f"GST Report ({GST_RATE}%)", [
                ("Gross revenue (incl. GST)", t["revenue"]), ("Taxable value", t["taxableValue"]),
                ("GST collected", t["gst"]),
            ])
        elif name == "series":
            ws.title = f"By {granularity.capitalize()}"
            series_sheet(ws, f"Breakdown by {granularity}")
        elif name == "transactions":
            ws.title = "Transactions"
            tx_sheet(ws)

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()


@router.get("/export")
def export(type: str = "all", granularity: str = "month", period: Optional[str] = "fy",
           frm: Optional[int] = Query(None, alias="from"), to: Optional[int] = None,
           admin: CurrentAdmin = Depends(require_superadmin)):
    a, b = _parse_window(period, frm, to)
    data = _xlsx(type.lower(), granularity, _rows(a, b), (a, b))
    fname = f"wimm_{type}_{granularity}_{datetime.fromtimestamp(a/1000, IST):%Y%m%d}.xlsx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
